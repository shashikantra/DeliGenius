from __future__ import annotations

from pathlib import Path
from typing import ClassVar

from openpyxl import load_workbook

from .base import ParsedDocument, ParsedElement, ParsedPage, register


@register
class XlsxParser:
    name: ClassVar[str] = "xlsx"
    supported_mimes: ClassVar[set[str]] = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    supported_extensions: ClassVar[set[str]] = {".xlsx", ".xlsm"}

    MAX_ROWS_PER_CHUNK_GROUP: ClassVar[int] = 200

    def parse(self, path: Path) -> ParsedDocument:
        wb = load_workbook(filename=path, read_only=True, data_only=True)
        try:
            pages: list[ParsedPage] = []
            sheet_names = [ws.title for ws in wb.worksheets]
            for page_num, ws in enumerate(wb.worksheets, start=1):
                elements: list[ParsedElement] = []
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    pages.append(ParsedPage(number=page_num, elements=[]))
                    continue
                header = ["" if c is None else str(c) for c in rows[0]]
                body = rows[1:]
                for start in range(0, len(body), self.MAX_ROWS_PER_CHUNK_GROUP):
                    group = body[start:start + self.MAX_ROWS_PER_CHUNK_GROUP]
                    table_rows = [header] + [
                        ["" if c is None else str(c) for c in r] for r in group
                    ]
                    elements.append(ParsedElement(
                        kind="table",
                        text=f"Sheet: {ws.title}\n" + "\n".join(" | ".join(r) for r in table_rows),
                        page=page_num,
                        rows=table_rows,
                        caption=f"Sheet {ws.title}",
                        meta={"sheet": ws.title,
                              "row_start": start + 2,
                              "row_end": start + 1 + len(group),
                              "header_repeated": True},
                    ))
                pages.append(ParsedPage(number=page_num, elements=elements))
        finally:
            wb.close()
        return ParsedDocument(
            source_path=str(path),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            pages=pages,
            metadata={"sheets": sheet_names},
            parser_name=self.name,
        )